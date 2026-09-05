"""
Generates a sample .docx architecture decision document under data/samples/
so reviewers can immediately test the Document Upload / Ingestion flow
without hunting for a real file.

Usage: python scripts/generate_sample_docx.py
"""
import os
import docx

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "samples")
os.makedirs(OUT_DIR, exist_ok=True)


def build(path, title, version_label, status, mechanism, extra_note):
    d = docx.Document()
    d.add_heading(title, level=0)
    d.add_paragraph(f"Version: {version_label}")
    d.add_paragraph(f"Status: {status}")
    d.add_paragraph(f"Owner: Security Architecture Team")
    d.add_heading("Section 1: Overview", level=1)
    d.add_paragraph(f"This document defines the authentication mechanism under evaluation. "
                     f"The system uses {mechanism} for user authentication across internal services.")
    d.add_heading("Section 2: Token Handling", level=1)
    d.add_paragraph("Access tokens are short-lived and refresh tokens are rotated on every use. "
                     "Tokens are signed and validated on every downstream request.")
    d.add_heading("Section 3: Multi-Factor Authentication", level=1)
    d.add_paragraph("MFA is required for all administrative accounts and for any account with "
                     "production write access.")
    d.add_heading("Section 4.2: Protocol Selection", level=1)
    d.add_paragraph(f"The approved authentication architecture uses {mechanism} as the primary "
                     f"protocol for identity verification across internal and external services. {extra_note}")
    d.add_heading("Section 5: Session Expiry", level=1)
    d.add_paragraph("Sessions expire after 30 minutes of inactivity and require re-authentication "
                     "for sensitive operations.")
    d.save(path)
    print("wrote", path)


if __name__ == "__main__":
    # A DRAFT v4 (newer, but should NOT become authoritative on its own)
    build(
        os.path.join(OUT_DIR, "Authentication_Architecture_v4_DRAFT.docx"),
        "Authentication Architecture", "4", "DRAFT",
        "OAuth 2.0 with OpenID Connect and passkey-based WebAuthn",
        "This revision is still under committee review and has not yet been approved.",
    )
    # An APPROVED v3 (older, but should win because it's approved)
    build(
        os.path.join(OUT_DIR, "Authentication_Architecture_v3_APPROVED.docx"),
        "Authentication Architecture", "3", "APPROVED",
        "OAuth 2.0 with OpenID Connect",
        "This decision was ratified by the Architecture Board on 2026-08-10.",
    )
